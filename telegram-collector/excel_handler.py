import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DEFAULT_HEADERS = [
    "Thời gian",
    "Loại tài khoản",
    "Gói / Thời hạn",
    "Giá bán (VNĐ)",
    "Bảo hành / Ghi chú",
    "Người bán",
    "Username / Liên hệ",
    "Tên nhóm",
    "Link tin nhắn",
    "Nội dung gốc"
]

class ExcelHandler:
    def __init__(self, file_path: str = "data.xlsx", sheet_name: str = "AccountSales"):
        self.file_path = file_path
        self.sheet_name = sheet_name
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """Khởi tạo file Excel nếu chưa tồn tại với header định dạng đẹp."""
        if not os.path.exists(self.file_path):
            wb = Workbook()
            ws = wb.active
            ws.title = self.sheet_name
            ws.append(DEFAULT_HEADERS)

            # Styling header
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            for col_num in range(1, len(DEFAULT_HEADERS) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Độ rộng cột
            col_widths = {
                "A": 20, # Thời gian
                "B": 22, # Loại tài khoản
                "C": 18, # Gói / Thời hạn
                "D": 18, # Giá bán (VNĐ)
                "E": 28, # Bảo hành / Ghi chú
                "F": 20, # Người bán
                "G": 22, # Username / Liên hệ
                "H": 25, # Tên nhóm
                "I": 35, # Link tin nhắn
                "J": 45  # Nội dung gốc
            }
            for col, width in col_widths.items():
                ws.column_dimensions[col].width = width

            wb.save(self.file_path)
            print(f"[ExcelHandler] Đã khởi tạo file Excel: {self.file_path}")

    def append_account_offer(self, account_type: str, plan: str, price: str,
                             warranty: str, seller_name: str, contact_info: str,
                             group_name: str, message_link: str, original_text: str,
                             timestamp: datetime = None) -> bool:
        """Ghi thêm một bản ghi bán tài khoản vào file Excel."""
        if timestamp is None:
            timestamp = datetime.now()

        row_data = [
            timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            account_type or "Chưa rõ",
            plan or "",
            price or "Thương lượng",
            warranty or "",
            seller_name or "",
            contact_info or "",
            group_name or "",
            message_link or "",
            original_text or ""
        ]

        try:
            wb = load_workbook(self.file_path)
            if self.sheet_name in wb.sheetnames:
                ws = wb[self.sheet_name]
            else:
                ws = wb.create_sheet(self.sheet_name)
                ws.append(DEFAULT_HEADERS)

            ws.append(row_data)

            # Căn lề
            current_row = ws.max_row
            ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=3).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=4).alignment = Alignment(horizontal="center")

            wb.save(self.file_path)
            print(f"  -> [Excel] +1 Dòng: [{account_type}] - {price} ({seller_name})")
            return True
        except PermissionError:
            print(f"[CẢNH BÁO] Không thể lưu file '{self.file_path}' vì file đang mở trong ứng dụng khác. Hãy đóng file để cập nhật!")
            return False
        except Exception as e:
            print(f"[LỖI Excel] {e}")
            return False
