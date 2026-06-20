import docx
import office2pdf
import os

# 1. Create a dummy docx
doc = docx.Document()
doc.add_heading('Test Document', 0)
doc.add_paragraph('Hello world! This is a test DOCX file to verify the PDF converter.')
doc.add_paragraph('It contains some basic formatting, headings, and text.')
doc.save('test.docx')

print("Created test.docx")

# 2. Convert to PDF
try:
    result = office2pdf.convert_path('test.docx')
    with open('test.pdf', 'wb') as f:
        f.write(result.pdf)
    print("Successfully converted test.docx to test.pdf!")
    if os.path.exists('test.pdf') and os.path.getsize('test.pdf') > 0:
        print(f"Verified test.pdf exists and is {os.path.getsize('test.pdf')} bytes.")
    else:
        print("Verification failed: test.pdf does not exist or is empty.")
except Exception as e:
    print(f"Conversion failed with error: {e}")
